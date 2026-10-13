"""
Données françaises du brief sur la taxe sur la valeur foncière (LVT).

Les fichiers bruts ne sont pas versionnés (DVF 2024 pèse 90 Mo). Les télécharger dans un
répertoire BRUT, puis :

    python extract_lvt.py BRUT

Fichiers bruts attendus dans BRUT :
  T_8201.xlsx, T_8210.xlsx  Insee, comptes de patrimoine annuels base 2020
                            https://www.insee.fr/fr/statistiques/8574720
  pib.xml                   Insee BDM, flux CNA-2020-PIB (SDMX)
                            https://bdm.insee.fr/series/sdmx/data/CNA-2020-PIB
  full2024.csv.gz           DVF géolocalisées 2024 (Etalab)
                            https://files.data.gouv.fr/geo-dvf/latest/csv/2024/full.csv.gz
  IA91.xlsx                 données de l'Insee Analyses n° 91 (André et Meslin, 2023)
                            https://www.insee.fr/fr/statistiques/7735076

Sorties dans donnees/ :
  patrimoine.csv       valeurs de fin d'année (Md€) et PIB (M€), 1978-2024
  dep2024.csv          prix médian des terrains à bâtir et des maisons par département
  tf_vs_prix.csv       taxe foncière / valeur des logements (2017) et prix au m² (2024) par commune
  tf_vs_prix_bins.csv  la même relation par vingtième de prix au m²
"""
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

BRUT = Path(sys.argv[1])
OUT = Path(__file__).resolve().parent / "donnees"
OUT.mkdir(exist_ok=True)


# ---------------------------------------------------------------- comptes de patrimoine
def pib_courant():
    t = (BRUT / "pib.xml").read_text(encoding="utf-8")
    cle = "Produit intérieur brut approche produit - Total des secteurs institutionnels - " \
          "Total des produits - Prix courant\""
    for attrs, corps in re.findall(r"<Series ([^>]*)>(.*?)</Series>", t, re.S):
        if cle in attrs:   # millions d'euros (UNIT_MULT = 6)
            return {int(a): float(v) for a, v in
                    re.findall(r'TIME_PERIOD="(\d+)" OBS_VALUE="([^"]*)"', corps)}
    raise ValueError("PIB introuvable")


def compte(nom):
    ws = openpyxl.load_workbook(BRUT / f"{nom}.xlsx", data_only=True)[nom]
    lignes = list(ws.iter_rows(values_only=True))
    annees = [int(x) for x in lignes[3][2:] if x]
    return {l[0]: {a: l[2 + i] for i, a in enumerate(annees)} for l in lignes[5:] if l[0]}


S1, S14 = compte("T_8201"), compte("T_8210")
pat = pd.DataFrame({
    "pib": pd.Series(pib_courant()),
    "terrains_batis": pd.Series(S1["N2111N"]), "logements": pd.Series(S1["N111N"]),
    "terrains_batis_men": pd.Series(S14["N2111N"]), "logements_men": pd.Series(S14["N111N"]),
    "terrains_cultives": pd.Series(S1["N2112N"]), "autres_terrains": pd.Series(S1["N2119N"]),
}).loc[1978:2024]
pat.to_csv(OUT / "patrimoine.csv")

# ---------------------------------------------------------------- DVF 2024
cols = ["id_mutation", "nature_mutation", "valeur_fonciere", "code_commune", "code_departement",
        "id_parcelle", "code_type_local", "surface_reelle_bati", "code_nature_culture",
        "surface_terrain"]
d = pd.read_csv(BRUT / "full2024.csv.gz", usecols=cols, low_memory=False,
                dtype={"code_commune": str, "code_departement": str, "code_type_local": str})
d = d[d.nature_mutation.isin(["Vente", "Vente terrain à bâtir"])]

# Terrains à bâtir : mutations dont toutes les lignes sont classées « terrain à bâtir » (AB) au
# cadastre, sans local ; entre 100 et 5 000 m².
g = d.groupby("id_mutation")
tab = g.agg(val=("valeur_fonciere", "first"), dep=("code_departement", "first"),
            n_local=("code_type_local", lambda s: s.notna().sum()),
            ab=("code_nature_culture", lambda s: (s == "AB").all()),
            surf=("surface_terrain", "sum"))
tab = tab[tab.ab & (tab.n_local == 0) & tab.surf.between(100, 5000) & (tab.val > 1000)].copy()
tab["pm2"] = tab.val / tab.surf

# Maisons : ventes d'une seule maison, sans appartement ni local d'activité.
v = d[d.nature_mutation == "Vente"]
loc = v[v.code_type_local.notna()]
k = loc.groupby("id_mutation").agg(
    nm=("code_type_local", lambda s: (s == "1").sum()),
    na=("code_type_local", lambda s: (s == "2").sum()),
    nl=("code_type_local", lambda s: (s == "4").sum()),
    sb=("surface_reelle_bati", lambda s: s[loc.loc[s.index, "code_type_local"] == "1"].max()))
ok = k[(k.nm == 1) & (k.na == 0) & (k.nl == 0)].index
m = v[v.id_mutation.isin(ok)]
st = m.drop_duplicates(["id_mutation", "id_parcelle", "surface_terrain"]) \
      .groupby("id_mutation").surface_terrain.sum()
mai = m.groupby("id_mutation").agg(val=("valeur_fonciere", "first"), dep=("code_departement", "first"))
mai["terrain"], mai["sb"] = st, k.loc[mai.index, "sb"]
mai = mai[(mai.val > 10000) & mai.sb.between(30, 400) & mai.terrain.between(50, 10000)]

dep = tab.groupby("dep").agg(n_t=("pm2", "size"), pm2_terrain=("pm2", "median"),
                             surf_tab=("surf", "median")).join(
      mai.groupby("dep").agg(n_m=("val", "size"), prix_maison=("val", "median"),
                             terrain_maison=("terrain", "median"), sb=("sb", "median")), how="outer")
# Part du terrain dans le prix d'une maison médiane : prix du m² de terrain à bâtir multiplié par
# la surface de parcelle médiane, rapporté au prix médian. Ordre de grandeur seulement.
dep["part_terrain"] = dep.pm2_terrain * dep.terrain_maison / dep.prix_maison
dep[dep.n_t >= 30].to_csv(OUT / "dep2024.csv")

# Prix médian au m² des logements par commune : ventes d'un seul logement (maison ou appartement).
k2 = loc.groupby("id_mutation").agg(
    n12=("code_type_local", lambda s: s.isin(["1", "2"]).sum()),
    n4=("code_type_local", lambda s: (s == "4").sum()),
    sb=("surface_reelle_bati", lambda s: s[loc.loc[s.index, "code_type_local"].isin(["1", "2"])].sum()),
    val=("valeur_fonciere", "first"), com=("code_commune", "first"))
k2 = k2[(k2.n12 == 1) & (k2.n4 == 0) & k2.sb.between(10, 400) & (k2.val > 5000)]
k2["pm2"] = k2.val / k2.sb
k2 = k2[k2.pm2.between(200, 40000)]
com = k2.groupby("com").agg(n=("pm2", "size"), pm2=("pm2", "median"))

# ---------------------------------------------------------------- taxe foncière / valeur (Insee)
ws = openpyxl.load_workbook(BRUT / "IA91.xlsx", data_only=True)["Figure 5a"]
tf = pd.DataFrame([r[:2] for r in ws.iter_rows(min_row=4, values_only=True)
                   if r[0] and isinstance(r[1], (int, float))], columns=["com", "tf"]).set_index("com")
tvp = tf.join(com, how="inner")
tvp = tvp[tvp.n >= 20].copy()
tvp["dec"] = pd.qcut(tvp.pm2, 20, labels=False)
tvp.to_csv(OUT / "tf_vs_prix.csv")
tvp.groupby("dec").agg(pm2=("pm2", "median"), tf=("tf", "median"),
                       tf25=("tf", lambda s: s.quantile(.25)), tf75=("tf", lambda s: s.quantile(.75)),
                       n=("tf", "size")).to_csv(OUT / "tf_vs_prix_bins.csv")
print("communes appariées :", len(tvp),
      "| élasticité log(TF/valeur) sur log(prix) :", round(np.polyfit(np.log(tvp.pm2), np.log(tvp.tf), 1)[0], 3))
